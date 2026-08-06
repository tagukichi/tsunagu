"""PPTX の各シェイプ位置・テキストを読み取り、HTMLに再現してQA用に描画する。
LibreOffice が使えない環境向けの簡易プレビュー。"""
import sys, html
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
SCALE = 96  # 1inch = 96px

def color_of(fmt, default=None):
    try:
        if fmt.type is not None and fmt.fore_color and fmt.fore_color.type is not None:
            return '#' + str(fmt.fore_color.rgb)
    except Exception:
        pass
    return default

def render(path, out):
    prs = Presentation(path)
    sw = prs.slide_width / EMU_IN
    sh = prs.slide_height / EMU_IN
    parts = ["""<meta charset="utf-8"><style>
    body{margin:0;background:#5a5a5a;font-family:'IPAPGothic','IPAGothic',sans-serif}
    .slide{position:relative;background:#fff;margin:18px auto;overflow:hidden;box-shadow:0 4px 18px rgba(0,0,0,.4)}
    .n{position:absolute;left:6px;top:4px;color:#bbb;font-size:11px;z-index:99}
    .sh{position:absolute;box-sizing:border-box}
    .tx{position:absolute;box-sizing:border-box;display:flex;white-space:pre-wrap;word-break:break-word;overflow:visible}
    </style>"""]
    for idx, slide in enumerate(prs.slides, 1):
        bg = '#FFFFFF'
        try:
            if slide.background.fill.type is not None:
                c = color_of(slide.background.fill)
                if c: bg = c
        except Exception:
            pass
        parts.append(f'<div class="slide" style="width:{sw*SCALE}px;height:{sh*SCALE}px;background:{bg}">')
        parts.append(f'<div class="n">slide {idx}</div>')
        for shp in slide.shapes:
            try:
                x = shp.left / EMU_IN * SCALE; y = shp.top / EMU_IN * SCALE
                w = shp.width / EMU_IN * SCALE; h = shp.height / EMU_IN * SCALE
            except Exception:
                continue
            st = shp.shape_type
            has_text = shp.has_text_frame and shp.text_frame.text.strip()
            # 図形（塗り・線）
            fill = None; line = None; radius = '0'
            try: fill = color_of(shp.fill)
            except Exception: pass
            try:
                if shp.line.color and shp.line.color.type is not None:
                    line = '#' + str(shp.line.color.rgb)
            except Exception: pass
            # 図形種別は prstGeom から取得
            nm = ''
            try:
                import re as _re
                m = _re.search(r'prstGeom prst="([^"]+)"', shp._element.xml)
                nm = (m.group(1) if m else '').lower()
            except Exception:
                nm = (shp.name or '').lower()
            style = f'left:{x}px;top:{y}px;width:{w}px;height:{h}px;'
            if 'ellipse' in nm: style += 'border-radius:50%;'
            elif 'round' in nm: style += 'border-radius:10px;'
            if 'triangle' in nm:
                style += f'background:transparent;border-left:{w/2}px solid transparent;border-right:{w/2}px solid transparent;border-bottom:{h}px solid {fill or "#999"};width:0;height:0;'
                parts.append(f'<div class="sh" style="{style}"></div>')
                continue
            if fill: style += f'background:{fill};'
            if line: style += f'border:1.5px solid {line};'
            if fill or line:
                parts.append(f'<div class="sh" style="{style}"></div>')
            if not has_text: continue
            # テキスト
            tf = shp.text_frame
            p0 = tf.paragraphs[0]
            algn = str(p0.alignment or '').lower()
            just = 'center' if 'center' in algn else ('flex-end' if 'right' in algn else 'flex-start')
            va = str(tf.vertical_anchor or '').lower()
            align_items = 'center' if 'middle' in va or 'center' in va else 'flex-start'
            runs = []
            for para in tf.paragraphs:
                seg = ''
                for r in para.runs:
                    sz = (r.font.size.pt if r.font.size else 14)
                    col = '#000000'
                    try:
                        if r.font.color and r.font.color.type is not None: col = '#' + str(r.font.color.rgb)
                    except Exception: pass
                    b = 'font-weight:700;' if r.font.bold else ''
                    seg += f'<span style="font-size:{sz}px;color:{col};{b}">{html.escape(r.text)}</span>'
                runs.append(seg or '&nbsp;')
            inner = '<br>'.join(runs)
            tstyle = (f'left:{x}px;top:{y}px;width:{w}px;height:{h}px;'
                      f'justify-content:{just};align-items:{align_items};line-height:1.25;')
            parts.append(f'<div class="tx" style="{tstyle}"><div style="width:100%;text-align:{ "center" if just=="center" else "left"}">{inner}</div></div>')
        parts.append('</div>')
    open(out, 'w').write('\n'.join(parts))
    print('wrote', out, 'slides:', len(prs.slides._sldIdLst))

render(sys.argv[1], sys.argv[2])
